#!/usr/bin/env python3
"""
fix_indentation.py — clean up text copied out of VS Code / a terminal.

Strips stray leading whitespace from every line and re-joins soft-wrapped
lines back into single paragraphs, so pasting into a .txt, .md, or .docx
file doesn't carry over ragged, indented line breaks. Headings, list items,
blockquotes, table rows, and fenced code blocks are each kept on their own
line rather than being merged together.

Usage:
    python fix_indentation.py <input> [output]

If <output> is omitted, the cleaned file is written next to this script
(the project root) as "<input-stem>_clean<input-extension>". A relative
<output> is also resolved against this script's folder, not your current
working directory.

Supported extensions on either side: .txt, .md, .docx
(.docx requires the optional `python-docx` package: pip install python-docx)
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent

FENCE_RE = re.compile(r"^(```|~~~)")
NEW_ITEM_RE = re.compile(r"^(#{1,6}\s|[-*+]\s|\d+[.)]\s|>\s?|\|)")
HR_RE = re.compile(r"^(-{3,}|\*{3,}|_{3,})\s*$")
TABLE_SEP_RE = re.compile(r"^\|?[\s:|-]+\|?$")


def _starts_new_item(line: str) -> bool:
    """Lines matching this stay on their own output line instead of being
    merged into the paragraph above (headings, bullets, numbered items,
    blockquote markers, table rows/separators, horizontal rules)."""
    stripped = line.strip()
    return bool(
        NEW_ITEM_RE.match(stripped)
        or HR_RE.match(stripped)
        or TABLE_SEP_RE.match(stripped)
    )


def _reflow_block(lines: list[str]) -> list[str]:
    """Collapse a blank-line-delimited block into its logical lines: a
    continuation line (no marker of its own) is dedented and appended to
    the item above it; anything that starts a new item stays separate."""
    items: list[str] = []
    current = ""
    started = False
    for line in lines:
        stripped = line.strip()
        if not started or _starts_new_item(line):
            if current:
                items.append(current)
            current = stripped
            started = True
        else:
            current = f"{current} {stripped}".strip()
    if current:
        items.append(current)
    return items


def clean_text(raw: str) -> str:
    lines = raw.splitlines()
    out_blocks: list[list[str]] = []
    block: list[str] = []
    in_fence = False
    fence_lines: list[str] = []

    def flush_block() -> None:
        nonlocal block
        if block:
            out_blocks.append(_reflow_block(block))
            block = []

    for line in lines:
        if in_fence:
            fence_lines.append(line)
            if FENCE_RE.match(line.strip()):
                out_blocks.append(fence_lines)
                fence_lines = []
                in_fence = False
            continue

        stripped = line.strip()
        if FENCE_RE.match(stripped):
            flush_block()
            in_fence = True
            fence_lines = [line]
            continue

        if stripped == "":
            flush_block()
            continue

        block.append(line)

    flush_block()
    if fence_lines:  # unterminated fence — emit what we have verbatim
        out_blocks.append(fence_lines)

    return "\n\n".join("\n".join(b) for b in out_blocks) + "\n"


def _require_docx():
    try:
        import docx  # type: ignore
    except ImportError as exc:
        raise SystemExit(
            "Reading/writing .docx needs the optional 'python-docx' package.\n"
            "Install it with:  pip install python-docx"
        ) from exc
    return docx


def read_input(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        docx = _require_docx()
        doc = docx.Document(str(path))
        return "\n\n".join(p.text for p in doc.paragraphs)
    return path.read_text(encoding="utf-8")


def write_output(path: Path, text: str) -> None:
    if path.suffix.lower() == ".docx":
        docx = _require_docx()
        doc = docx.Document()
        blocks = text.strip("\n").split("\n\n")
        for i, block in enumerate(blocks):
            for line in block.split("\n"):
                doc.add_paragraph(line)
            if i != len(blocks) - 1:
                doc.add_paragraph("")  # blank line between paragraphs/blocks
        doc.save(str(path))
    else:
        path.write_text(text, encoding="utf-8")


def main(argv: list[str]) -> int:
    if not argv or argv[0] in ("-h", "--help"):
        print(__doc__)
        return 0 if argv else 1

    input_path = Path(argv[0]).expanduser().resolve()
    if not input_path.exists():
        print(f"Input file not found: {input_path}", file=sys.stderr)
        return 1

    if len(argv) > 1:
        output_path = Path(argv[1]).expanduser()
        if not output_path.is_absolute():
            output_path = SCRIPT_DIR / output_path
    else:
        output_path = SCRIPT_DIR / f"{input_path.stem}_clean{input_path.suffix}"

    cleaned = clean_text(read_input(input_path))
    write_output(output_path, cleaned)
    print(f"Wrote: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
